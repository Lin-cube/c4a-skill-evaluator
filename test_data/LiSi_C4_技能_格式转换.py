#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""format_convert.py — csv/xlsx 互转。

用法：python format_convert.py 数据.csv
"""

import sys
from pathlib import Path


def main(path: str) -> int:
    p = Path(path)
    if p.suffix == ".csv":
        # 转 xlsx（需要 openpyxl）
        import openpyxl  # noqa: F401
        wb = openpyxl.Workbook()
        ws = wb.active
        for row in p.read_text(encoding="utf-8").splitlines():
            ws.append(row.split(","))
        wb.save(str(p.with_suffix(".xlsx")))
    elif p.suffix == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(p)
        ws = wb.active
        out = "\n".join(",".join(str(c) for c in r) for r in ws.iter_rows(values_only=True))
        p.with_suffix(".csv").write_text(out, encoding="utf-8")
    print("转换完成")
    return 0


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("用法：python format_convert.py 文件.csv|xlsx")
        sys.exit(1)
    sys.exit(main(sys.argv[1]))
